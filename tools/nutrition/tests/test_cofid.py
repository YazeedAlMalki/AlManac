import re
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path

import openpyxl

from almanac_nutrition.canonical_io import CanonicalError, read_canonical
from almanac_nutrition.lake import InputVerificationError, Lake, sha256_file
from almanac_nutrition.sources import cofid, load_sources
from almanac_nutrition.xlsx import sheet_file
from tests.support import DICTIONARY, REAL_LAKE, integration, register, temp_lake

LOCAL = "raw/cofid/2021"
HEADINGS = ["Food Code", "Food Name", "Description", "Group", "Previous", "Main data references",
            "Footnote", "Water (g)", "Protein (g)", "Fat (g)", "Carbohydrate (g)", "Energy (kcal) (kcal)",
            "Starch (g)", "Alcohol (g)", "NSP (g)", "AOAC fibre (g)"]
CODES = [None] * 7 + ["WATER", "PROT", "FAT", "CHO", "KCALS", "STAR", "ALCO", "ENGFIB", "AOACFIB"]
DESCRIPTIONS = [None] * 7 + ["Water", "Protein", "Fat", "Carbohydrate", "kcal", "Starch", "Alcohol",
                             "Non-starch polysaccharide", "AOAC fibre"]
# Spreadsheet rows 4-8. Text, float and int cells, as in the release.
FOODS = [
    ["13-145", "Ackee, canned, drained", "8 cans", "DG", 554, "MW4", None,
     "76.7", "2.9", "15.2", "0.8", "151", "Tr", None, "N", None],
    ["17-752", "Wine, red", "", "QE", None, "ref", None, "88.4", "0.1", "0", "0.2", "76", "0", 10.7, "0", "0"],
    ["13-669", "Aubergine, flesh and skin, roasted in rapeseed oil", "", "DG", None, "ref", None,
     83.9, 2.1, 3.8, 5.2, 62, "(0.5)", None, "3.1", "4.0"],
    ["13-669", "Watercress, raw", "", "DG", "13-653", "ref", None, 94.8, 1.9, 0.3, "Tr", 10, "0", None, "1.5", "N"],
    ["19-100", "Meat pie, fixture", "", "MR", None, "ref", None, "50", "10", "15", "N", "N", None, None, "N", None],
]


def fixture_lake(case, foods=FOODS, headings=HEADINGS):
    lake = temp_lake(case)
    base = lake.root / LOCAL
    base.mkdir(parents=True)
    book = openpyxl.Workbook()
    book.active.title = "List of tables"
    book.active.append(["List of tables in CoFID 2021"])
    sheet = book.create_sheet(cofid.FOOD_SHEET)
    for row in (headings, CODES, DESCRIPTIONS, *foods):
        sheet.append(row)
    book.create_sheet("1.4 Inorganics").append(["Food Code", "Sodium (mg)"])
    book.save(base / cofid.WORKBOOK)
    old = openpyxl.Workbook()
    old.active.title = "Old foods"
    old.active.append(["Food Code", "Southgate fibre (g)"])
    old.save(base / cofid.OLD_FOODS)
    register(lake, "cofid", LOCAL, list(cofid.INPUTS), release="2021")
    return lake


def canonical(lake):
    foods, names, values, manifest = read_canonical(lake.canonical("cofid"))
    return foods, names, {(v.food_ref, v.nutrient_id): v for v in values}, manifest


class Extract(unittest.TestCase):
    def test_refuses_inputs_that_differ_from_the_collection_manifest(self):
        lake = fixture_lake(self)
        (lake.root / LOCAL / cofid.OLD_FOODS).write_bytes(b"tampered")
        with self.assertRaises(InputVerificationError):
            cofid.extract(lake)
        self.assertFalse(lake.extracted("cofid").exists())

    def test_every_sheet_of_both_workbooks_with_all_header_rows(self):
        out = cofid.extract(fixture_lake(self))
        self.assertEqual(sorted(p.name for p in out.iterdir() if p.suffix == ".csv"), sorted([
            sheet_file(cofid.WORKBOOK, "List of tables"), sheet_file(cofid.WORKBOOK, cofid.FOOD_SHEET),
            sheet_file(cofid.WORKBOOK, "1.4 Inorganics"), sheet_file(cofid.OLD_FOODS, "Old foods")]))
        lines = (out / sheet_file(cofid.WORKBOOK, cofid.FOOD_SHEET)).read_text().splitlines()
        self.assertEqual(len(lines), 3 + len(FOODS))
        self.assertIn('"PROT"', lines[1])


class Canonicalise(unittest.TestCase):
    def setUp(self):
        self.lake = fixture_lake(self)
        cofid.extract(self.lake)
        cofid.canonicalise(self.lake, DICTIONARY)
        self.foods, self.names, self.values, self.manifest = canonical(self.lake)

    def test_tokens_have_no_amount_and_blanks_have_no_row(self):
        watercress = "cofid:13-669@row7"
        carbohydrate = self.values[(watercress, "carbohydrate_available_monosaccharide")]
        self.assertEqual((carbohydrate.amount, carbohydrate.qualifier, carbohydrate.source_value),
                         (None, "trace", "Tr"))
        fibre = self.values[(watercress, "fibre_total_dietary")]
        self.assertEqual((fibre.amount, fibre.qualifier, fibre.source_value), (None, "not_analysed", "N"))
        self.assertNotIn(("cofid:13-145", "fibre_total_dietary"), self.values)
        self.assertNotIn(("cofid:13-145", "alcohol"), self.values)
        energy = self.values[("cofid:19-100", "energy_kcal")]
        self.assertEqual((energy.amount, energy.qualifier), (None, "not_analysed"))

    def test_numbers_zeros_and_units(self):
        ackee = self.values[("cofid:13-145", "carbohydrate_available_monosaccharide")]
        self.assertEqual((ackee.amount, ackee.qualifier, ackee.source_nutrient_id), (0.8, "measured", "CHO"))
        energy = self.values[("cofid:13-145", "energy_kcal")]
        self.assertEqual((energy.amount, energy.qualifier, energy.source_unit), (151.0, "calculated_factor", "kcal"))
        fat = self.values[("cofid:17-752", "fat_total")]
        self.assertEqual((fat.amount, fat.qualifier, fat.source_value), (0.0, "zero_reported", "0"))
        alcohol = self.values[("cofid:17-752", "alcohol")]
        self.assertEqual((alcohol.amount, alcohol.source_value), (10.7, "10.7"))  # a float cell, verbatim

    def test_alcoholic_beverages_are_per_100ml_and_nothing_else_is(self):
        bases = {(ref, v.basis) for (ref, _), v in self.values.items()}
        self.assertEqual({basis for ref, basis in bases if ref == "cofid:17-752"}, {"per_100ml"})
        self.assertEqual({basis for ref, basis in bases if ref != "cofid:17-752"}, {"per_100g"})

    def test_a_repeated_food_code_keeps_both_foods(self):
        self.assertEqual([f.food_ref for f in self.foods if f.local_id.startswith("13-669")],
                         ["cofid:13-669@row6", "cofid:13-669@row7"])
        self.assertEqual(self.manifest["notes"]["collisions"][0]["names"],
                         ["Aubergine, flesh and skin, roasted in rapeseed oil", "Watercress, raw"])
        # "(0.5)" sits in Starch, which is not mapped, so it is extracted and left alone.
        self.assertNotIn(("cofid:13-669@row6", "starch"), self.values)

    def test_group_names_come_from_appendix_b(self):
        groups = {f.food_ref: (f.food_group_code, f.food_group_name) for f in self.foods}
        self.assertEqual(groups["cofid:17-752"], ("QE", "Wines"))
        self.assertEqual(groups["cofid:19-100"], ("MR", "Meat dishes"))

    def test_reruns_are_byte_identical_and_the_cli_finds_the_module(self):
        first = self.manifest["outputs"]
        cofid.canonicalise(self.lake, DICTIONARY)
        self.assertEqual(canonical(self.lake)[3]["outputs"], first)
        self.assertIs(load_sources()["cofid"], cofid)


class Refusals(unittest.TestCase):
    def run_both(self, lake):
        cofid.extract(lake)
        return cofid.canonicalise(lake, DICTIONARY)

    def test_a_bracketed_value_in_a_mapped_column_fails(self):
        foods = [FOODS[0][:8] + ["(2.9)"] + FOODS[0][9:]] + FOODS[1:]
        with self.assertRaisesRegex(CanonicalError, r"'\(2\.9\)' is neither a number"):
            self.run_both(fixture_lake(self, foods=foods))

    def test_a_unit_that_disagrees_with_the_mapping_fails(self):
        headings = [h.replace("Fat (g)", "Fat (mg)") for h in HEADINGS]
        with self.assertRaisesRegex(CanonicalError, "says 'Fat \\(mg\\)'"):
            self.run_both(fixture_lake(self, headings=headings))

    def test_an_extracted_sheet_edited_after_extraction_is_refused(self):
        lake = fixture_lake(self)
        out = cofid.extract(lake)
        sheet = out / sheet_file(cofid.WORKBOOK, cofid.FOOD_SHEET)
        sheet.write_text(sheet.read_text().replace('"Tr"', '"0"'))
        with self.assertRaisesRegex(CanonicalError, "changed after extraction"):
            cofid.canonicalise(lake, DICTIONARY)


def appendix_b(guide: Path) -> dict[str, str]:
    text = subprocess.run(["pdftotext", "-layout", str(guide), "-"], capture_output=True, text=True,
                          check=True).stdout
    section = text.split("Appendix B: Food sub-group codes")[-1].split("Appendix C")[0]
    return {m.group(2): m.group(1) for line in section.splitlines()
            if (m := re.fullmatch(r"\s*(\S.*?)\s{2,}([A-Z]{1,3})\s*", line))}


class AgainstTheRealLake(unittest.TestCase):
    @integration
    def test_every_food_both_collisions_and_the_alcoholic_beverages(self):
        with tempfile.TemporaryDirectory() as out:
            scratch = Lake(out)
            (Path(out) / "manifests").symlink_to(REAL_LAKE.manifests)
            (Path(out) / "raw").symlink_to(REAL_LAKE.root / "raw")
            cofid.extract(scratch)
            manifest = cofid.canonicalise(scratch, DICTIONARY)
            _, _, values, _ = read_canonical(scratch.canonical("cofid"))
        self.assertEqual(manifest["counts"]["foods"], 2887)
        self.assertEqual(manifest["notes"]["collisions"][0]["local_ids"],
                         ["13-669@row55", "13-669@row2827"])
        self.assertEqual(len(manifest["notes"]["per_100ml_foods"]), 31)
        self.assertFalse([v for v in values if v.amount == 0 and v.source_value.strip() in ("Tr", "N")])

    @integration
    @unittest.skipUnless(shutil.which("pdftotext"), "pdftotext is not installed")
    def test_the_group_name_table_is_the_user_guides_appendix_b(self):
        guide = REAL_LAKE.root / LOCAL / cofid.GUIDE
        self.assertEqual(sha256_file(guide), cofid.GUIDE_SHA256)
        self.assertEqual(appendix_b(guide), cofid.APPENDIX_B)


if __name__ == "__main__":
    unittest.main()
