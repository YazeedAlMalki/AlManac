import tempfile
import unittest
from pathlib import Path

import openpyxl

from almanac_nutrition.canonical_io import CanonicalError, read_canonical
from almanac_nutrition.lake import InputVerificationError, Lake
from almanac_nutrition.sources import afcd, load_sources
from almanac_nutrition.values import KJ_PER_KCAL
from tests.support import DICTIONARY, REAL_LAKE, integration, register, temp_lake

LOCAL = "raw/ausnut/release-3"
NUTRIENTS = ["Energy with dietary fibre, equated \n(kJ)", "Energy, without dietary fibre, equated \n(kJ)",
             "Moisture (water) \n(g)", "Protein \n(g)", "Fat, total \n(g)", "Total dietary fibre \n(g)",
             "Alcohol \n(g)", "Available carbohydrate, without sugar alcohols \n(g)",
             "Available carbohydrate, with sugar alcohols \n(g)"]
# The mL sheet spells some headings with different whitespace; the pipeline must not care.
NUTRIENTS_ML = [h.replace(" \n", "\n") for h in NUTRIENTS]
GRAMS = [
    # Cardamom's two carbohydrate columns differ so that the tests can tell which one is read.
    ["F002258", 31302, "Borrowed", "Cardamom seed, dried, ground", 1236, 1012, 8.3, 10.8, 6.7, 28, 0, 34.4, 36.0],
    ["F008936", 31304, "Recipe", "Stock, liquid, prepared from cube ", 18, 18, 100.1, 0.2, 0.2, 0, 0, 0.4, 0.4],
]
MILLILITRES = [["F008936", 31304, "Recipe", "Stock, liquid, prepared from cube", 19, 19, 105.0, 0.21, 0.21,
                0, 0, 0.42, 0.42]]


def fixture_lake(case, grams=GRAMS, millilitres=MILLILITRES):
    lake = temp_lake(case)
    base = lake.root / LOCAL
    base.mkdir(parents=True)
    profiles = openpyxl.Workbook()
    profiles.active.title = "Contents"
    profiles.active.append(["Release 3 - Nutrient profiles - Content summary"])
    for title, headings, rows in ((afcd.PER_100G, NUTRIENTS, grams), (afcd.PER_100ML, NUTRIENTS_ML, millilitres)):
        sheet = profiles.create_sheet(title)
        sheet.append([f"Release 3 - Nutrient profiles ({title})"])
        sheet.append(["(row 2 is blank in the release)"])
        sheet.append(["Public Food Key", "Classification", "Derivation", "Food Name", *headings])
        for row in rows:
            sheet.append(row)
    profiles.save(base / afcd.PROFILES)
    groups = openpyxl.Workbook()
    groups.active.title = afcd.GROUP_SHEET
    for row in (["Release 3 - Food group information"], ["General note"],
                ["Food group ID", "Food group name", "Inclusions", "Number of foods in Release 3"],
                [31, "Savoury sauces and condiments", "311 - Gravies and savoury sauces ", 90],
                [None, None, "313 - Herbs, spices, seasonings and stock cubes", None]):
        groups.active.append(row)
    groups.save(base / afcd.GROUPS)
    for workbook in afcd.WORKBOOKS:
        if workbook not in (afcd.PROFILES, afcd.GROUPS):
            other = openpyxl.Workbook()
            other.active.append(["Contents", workbook])
            other.save(base / workbook)
    register(lake, "ausnut", LOCAL, list(afcd.WORKBOOKS), release="release-3")
    return lake


def canonical(lake):
    foods, names, values, manifest = read_canonical(lake.canonical("afcd"))
    return foods, names, {(v.food_ref, v.nutrient_id, v.basis): v for v in values}, manifest


class Extract(unittest.TestCase):
    def test_refuses_inputs_that_differ_from_the_collection_manifest(self):
        lake = fixture_lake(self)
        (lake.root / LOCAL / afcd.PROFILES).write_bytes(b"tampered")
        with self.assertRaises(InputVerificationError):
            afcd.extract(lake)
        self.assertFalse(lake.extracted("afcd").exists())

    def test_every_sheet_of_every_workbook(self):
        out = afcd.extract(fixture_lake(self))
        self.assertEqual(len([p for p in out.iterdir() if p.suffix == ".csv"]), 8)


class Canonicalise(unittest.TestCase):
    def setUp(self):
        self.lake = fixture_lake(self)
        afcd.extract(self.lake)
        afcd.canonicalise(self.lake, DICTIONARY)
        self.foods, self.names, self.values, self.manifest = canonical(self.lake)

    def test_energy_is_converted_from_kilojoules_and_keeps_its_source_cell(self):
        energy = self.values[("afcd:F002258", "energy_kcal", "per_100g")]
        self.assertAlmostEqual(energy.amount, 1236 / KJ_PER_KCAL)
        self.assertEqual((energy.source_value, energy.source_unit, energy.qualifier),
                         ("1236", "kJ", "calculated_factor"))

    def test_the_food_level_derivation_sets_qualifier_and_confidence(self):
        protein = self.values[("afcd:F002258", "protein", "per_100g")]
        self.assertEqual((protein.amount, protein.qualifier, protein.confidence), (10.8, "borrowed", "Borrowed"))
        stock = self.values[("afcd:F008936", "protein", "per_100g")]
        self.assertEqual((stock.qualifier, stock.confidence), ("calculated_recipe", "Recipe"))
        alcohol = self.values[("afcd:F002258", "alcohol", "per_100g")]
        self.assertEqual((alcohol.amount, alcohol.qualifier), (0.0, "zero_reported"))

    def test_available_carbohydrate_includes_sugar_alcohols(self):
        carbohydrate = self.values[("afcd:F002258", "carbohydrate_available", "per_100g")]
        self.assertEqual((carbohydrate.amount, carbohydrate.source_nutrient_id),
                         (36.0, "Available carbohydrate, with sugar alcohols (g)"))

    def test_liquids_have_both_bases_and_they_are_not_merged(self):
        grams = self.values[("afcd:F008936", "carbohydrate_available", "per_100g")]
        millilitres = self.values[("afcd:F008936", "carbohydrate_available", "per_100ml")]
        self.assertEqual((grams.amount, millilitres.amount), (0.4, 0.42))
        self.assertNotIn(("afcd:F002258", "protein", "per_100ml"), self.values)
        self.assertEqual(self.manifest["notes"]["per_100ml_foods"], 1)
        self.assertEqual(len(self.values), 6 * 2 + 6)

    def test_names_groups_and_identifiers(self):
        self.assertEqual([f.food_ref for f in self.foods], ["afcd:F002258", "afcd:F008936"])
        self.assertEqual((self.foods[0].food_group_code, self.foods[0].food_group_name),
                         ("31302", "Herbs, spices, seasonings and stock cubes"))
        self.assertEqual(self.names[1].name, "Stock, liquid, prepared from cube")

    def test_reruns_are_byte_identical_and_the_cli_finds_the_module(self):
        first = self.manifest["outputs"]
        afcd.canonicalise(self.lake, DICTIONARY)
        self.assertEqual(canonical(self.lake)[3]["outputs"], first)
        self.assertIs(load_sources()["afcd"], afcd)


class Refusals(unittest.TestCase):
    def run_both(self, lake):
        afcd.extract(lake)
        return afcd.canonicalise(lake, DICTIONARY)

    def test_an_unmapped_derivation_fails(self):
        grams = [GRAMS[0][:2] + ["Guessed"] + GRAMS[0][3:], GRAMS[1]]
        with self.assertRaisesRegex(CanonicalError, "'Guessed' has no row"):
            self.run_both(fixture_lake(self, grams=grams))

    def test_a_liquid_missing_from_the_grams_sheet_fails(self):
        with self.assertRaisesRegex(CanonicalError, "does not name exactly one food"):
            self.run_both(fixture_lake(self, millilitres=[["F999999"] + MILLILITRES[0][1:]]))

    def test_a_token_in_a_mapped_cell_fails(self):
        grams = [GRAMS[0][:7] + ["Tr"] + GRAMS[0][8:], GRAMS[1]]
        with self.assertRaisesRegex(CanonicalError, "neither a number nor a known token"):
            self.run_both(fixture_lake(self, grams=grams))


class AgainstTheRealLake(unittest.TestCase):
    @integration
    def test_every_food_on_both_bases(self):
        with tempfile.TemporaryDirectory() as out:
            scratch = Lake(out)
            (Path(out) / "manifests").symlink_to(REAL_LAKE.manifests)
            (Path(out) / "raw").symlink_to(REAL_LAKE.root / "raw")
            afcd.extract(scratch)
            manifest = afcd.canonicalise(scratch, DICTIONARY)
        self.assertEqual(manifest["counts"]["foods"], 1588)
        self.assertEqual(manifest["notes"]["per_100ml_foods"], 213)
        self.assertEqual(manifest["counts"]["foods_without_values"], 0)


if __name__ == "__main__":
    unittest.main()
